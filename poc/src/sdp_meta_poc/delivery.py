"""Synthetic ZIP/Excel deliveries normalized before either pipeline reads them."""

import hashlib
import io
import zipfile
from pathlib import Path

import polars as pl
from openpyxl import Workbook, load_workbook

from .cases import SCHEMA, batch


def read_zip(payload):
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        entries = archive.infolist()
        if len(entries) != 1 or entries[0].filename != "events.csv":
            raise ValueError("fixture archive must contain exactly events.csv")
        if entries[0].file_size > 65536:
            raise ValueError("fixture archive exceeds size limit")
        return pl.read_csv(io.BytesIO(archive.read("events.csv")), schema=SCHEMA)


def read_excel(payload):
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    try:
        rows = list(workbook["events"].values)
        if tuple(rows[0]) != tuple(SCHEMA):
            raise ValueError("unexpected fixture workbook columns")
        if any(
            isinstance(value, str) and value.startswith("=") for row in rows[1:] for value in row
        ):
            raise ValueError("formulas are not accepted in fixture inputs")
        return pl.DataFrame(rows[1:], schema=SCHEMA, orient="row")
    finally:
        workbook.close()


def make_delivery(name, stage, root: Path):
    frame = batch(name)
    buffer = io.BytesIO()
    if name == "initial":
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("events.csv", frame.write_csv())
        payload, extension = buffer.getvalue(), "zip"
        normalized = read_zip(payload)
    elif name == "incremental":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "events"
        sheet.append(list(SCHEMA))
        for row in frame.rows():
            sheet.append(row)
        workbook.save(buffer)
        workbook.close()
        payload, extension = buffer.getvalue(), "xlsx"
        normalized = read_excel(payload)
    else:
        payload, extension = frame.write_csv().encode(), "csv"
        normalized = pl.read_csv(io.BytesIO(payload), schema=SCHEMA)
    original = root / "original" / f"{stage}.{extension}"
    target = root / "landing" / f"{stage}.csv"
    original.parent.mkdir(parents=True, exist_ok=True)
    target.parent.mkdir(parents=True, exist_ok=True)
    if original.exists() or target.exists():
        raise FileExistsError(f"delivery {stage} already exists")
    original.write_bytes(payload)
    normalized.write_csv(target)
    return normalized, {
        "original": str(original),
        "original_sha256": hashlib.sha256(payload).hexdigest(),
        "normalized": str(target),
        "normalized_sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
        "rows": normalized.height,
    }
